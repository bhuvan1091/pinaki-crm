from dotenv import load_dotenv
from pathlib import Path
load_dotenv(Path(__file__).parent / ".env", override=True)

import os
import io
import csv
import logging
import secrets
import asyncio
import base64
from io import BytesIO
from datetime import datetime, timezone, timedelta
from typing import Optional, List

import bcrypt
import jwt
import resend
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfgen import canvas as pdfcanvas
from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from motor.motor_asyncio import AsyncIOMotorClient, AsyncIOMotorGridFSBucket
from bson import ObjectId
from pydantic import BaseModel, Field, EmailStr

# ---------- setup ----------
ROOT_DIR = Path(__file__).parent
mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]
gridfs = AsyncIOMotorGridFSBucket(db)
app = FastAPI(title="Pinaki Solutions CRM")
api = APIRouter(prefix="/api")

JWT_ALGORITHM = "HS256"
ROLES = ["admin", "sales", "design", "production", "dispatch", "accounts", "management"]
STAGES = ["Order Received", "Design", "Client Approval", "Production", "Dispatch", "Challan", "Delivery", "Accounts", "Invoice", "Payment"]
LEAD_STATUSES = ["New", "Contacted", "Qualified", "Proposal Sent", "Negotiation", "Won", "Lost"]

resend.api_key = os.environ.get("RESEND_API_KEY", "")
SENDER_EMAIL = os.environ.get("SENDER_EMAIL", "info@ashokatechnovations.com")

def frontend_url(request=None):
    env_url = os.environ.get("FRONTEND_URL")
    if env_url:
        return env_url.rstrip("/")
    if request is not None:
        origin = request.headers.get("origin")
        if origin:
            return origin.rstrip("/")
        host = request.headers.get("host")
        if host:
            return f"https://{host}".rstrip("/")
    return ""

def sign_approval_token(design_id, order_id):
    return jwt.encode(
        {"sub": design_id, "order_id": order_id, "type": "approval_link",
         "exp": datetime.now(timezone.utc) + timedelta(days=14)},
        secret(), algorithm=JWT_ALGORITHM,
    )

def verify_approval_token(token):
    payload = jwt.decode(token, secret(), algorithms=[JWT_ALGORITHM])
    if payload.get("type") != "approval_link":
        raise jwt.InvalidTokenError("Not an approval link")
    return payload

def sign_pay_token(invoice_id):
    return jwt.encode(
        {"sub": invoice_id, "type": "pay_link", "exp": datetime.now(timezone.utc) + timedelta(days=90)},
        secret(), algorithm=JWT_ALGORITHM,
    )

def verify_pay_token(token):
    payload = jwt.decode(token, secret(), algorithms=[JWT_ALGORITHM])
    if payload.get("type") != "pay_link":
        raise jwt.InvalidTokenError("Not a payment link")
    return payload

BANK_DETAILS = {
    "account_name": "Pinaki Solutions Pvt. Ltd.",
    "bank_name": "HDFC Bank",
    "account_number": "50200012345678",
    "ifsc": "HDFC0001234",
    "branch": "Mumbai — Fort",
    "upi_id": "pinakisolutions@hdfc",
    "swift": "HDFCINBB",
}

def rupees(n):
    return f"Rs. {float(n or 0):,.2f}"

def generate_invoice_pdf(order, invoice, client):
    return _render_document_pdf(order, client, "Tax Invoice", f"Invoice #: {invoice['invoice_number']}", f"Date: {invoice['invoice_date']}", extra_meta=[("PO Number", order.get("po_number") or "—"), ("Payment Terms", invoice.get("payment_terms", "")), ("Due Date", invoice.get("due_date", ""))], show_totals=True, invoice=invoice)

def generate_challan_pdf(order, challan, client):
    return _render_document_pdf(order, client, "Delivery Challan", f"Challan #: {challan['challan_number']}", f"Date: {challan['challan_date']}", extra_meta=[("Order ID", order.get("order_id", "")), ("PO Number", order.get("po_number") or "—"), ("Transporter", challan.get("transporter") or "—"), ("Tracking / LR", challan.get("tracking_number") or "—")], show_totals=False, quantity_override=challan.get("quantity"))

def _render_document_pdf(order, client, title, ref_top, ref_bottom, extra_meta, show_totals, invoice=None, quantity_override=None):
    buf = BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A4)
    W, H = A4
    c.setFillColor(colors.HexColor("#0f172a"))
    c.rect(0, H - 30 * mm, W, 30 * mm, stroke=0, fill=1)
    c.setFillColor(colors.HexColor("#93c5fd"))
    c.setFont("Helvetica-Bold", 9)
    c.drawString(20 * mm, H - 12 * mm, "PINAKI SOLUTIONS")
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 22)
    c.drawString(20 * mm, H - 22 * mm, title)
    c.setFillColor(colors.HexColor("#cbd5e1"))
    c.setFont("Helvetica", 9)
    c.drawRightString(W - 20 * mm, H - 16 * mm, ref_top)
    c.drawRightString(W - 20 * mm, H - 22 * mm, ref_bottom)

    y = H - 45 * mm
    c.setFillColor(colors.HexColor("#64748b"))
    c.setFont("Helvetica", 8)
    c.drawString(20 * mm, y, "FROM")
    c.drawString(110 * mm, y, "TO")
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica-Bold", 11)
    c.drawString(20 * mm, y - 6 * mm, "Pinaki Solutions Pvt. Ltd.")
    c.drawString(110 * mm, y - 6 * mm, order.get("client_name", ""))
    c.setFont("Helvetica", 9)
    c.setFillColor(colors.HexColor("#334155"))
    c.drawString(20 * mm, y - 12 * mm, "info@ashokatechnovations.com")
    c.drawString(20 * mm, y - 17 * mm, "GSTIN: 27AAECP1234N1ZQ")
    if client:
        c.drawString(110 * mm, y - 12 * mm, (client.get("shipping_address") or client.get("billing_address") or "")[:60])
        c.drawString(110 * mm, y - 17 * mm, f"GSTIN: {client.get('gstin') or '—'}")

    y = H - 75 * mm
    for label, val in extra_meta:
        c.setFillColor(colors.HexColor("#94a3b8"))
        c.setFont("Helvetica", 8)
        c.drawString(20 * mm, y, str(label).upper())
        c.setFillColor(colors.HexColor("#0f172a"))
        c.setFont("Helvetica-Bold", 10)
        c.drawString(20 * mm, y - 5 * mm, str(val))
        y -= 12 * mm

    # line item
    y = H - 130 * mm
    c.setFillColor(colors.HexColor("#0f172a"))
    c.rect(20 * mm, y, W - 40 * mm, 8 * mm, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(23 * mm, y + 2.5 * mm, "DESCRIPTION")
    c.drawRightString(120 * mm, y + 2.5 * mm, "QTY")
    if show_totals:
        c.drawRightString(150 * mm, y + 2.5 * mm, "RATE")
        c.drawRightString(W - 23 * mm, y + 2.5 * mm, "AMOUNT")
    else:
        c.drawRightString(W - 23 * mm, y + 2.5 * mm, "UOM")

    y -= 10 * mm
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica", 10)
    c.drawString(23 * mm, y, order.get("product", ""))
    c.setFillColor(colors.HexColor("#64748b"))
    c.setFont("Helvetica", 8)
    c.drawString(23 * mm, y - 4 * mm, order.get("product_code") or "")
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica", 10)
    qty = quantity_override if quantity_override is not None else order.get("quantity", 0)
    c.drawRightString(120 * mm, y, f"{qty:,}")
    if show_totals:
        c.drawRightString(150 * mm, y, rupees(order.get("unit_price", 0)))
        c.drawRightString(W - 23 * mm, y, rupees(order.get("order_value", 0)))
    else:
        c.drawRightString(W - 23 * mm, y, "units")

    if show_totals and invoice:
        y -= 25 * mm
        for label, val, bold in [("Subtotal", order.get("order_value", 0), False),
                                 ("GST (18%)", order.get("gst", 0), False),
                                 ("TOTAL", invoice.get("total", 0), True)]:
            c.setFillColor(colors.HexColor("#0f172a") if bold else colors.HexColor("#475569"))
            c.setFont("Helvetica-Bold" if bold else "Helvetica", 11 if bold else 10)
            c.drawRightString(150 * mm, y, label)
            c.drawRightString(W - 23 * mm, y, rupees(val))
            y -= 7 * mm

    # footer
    c.setStrokeColor(colors.HexColor("#e2e8f0"))
    c.line(20 * mm, 32 * mm, W - 20 * mm, 32 * mm)
    c.setFillColor(colors.HexColor("#94a3b8"))
    c.setFont("Helvetica", 8)
    if show_totals:
        c.drawString(20 * mm, 24 * mm, "Thank you for your business. Please make payment by the due date shown above.")
    else:
        c.drawString(20 * mm, 24 * mm, "Please verify the goods against this challan. Report any discrepancy within 24 hours.")
        c.drawString(20 * mm, 19 * mm, "Received in good condition. Signature: __________________________________")
    c.drawString(20 * mm, 14 * mm, "This is a computer-generated document and does not require a signature.")
    c.showPage()
    c.save()
    return buf.getvalue()

EMAIL_TEMPLATES = {
    "design_share": {"subject": "Design ready for your review — Order {order_id}", "intro": "Please find attached the latest design for your approval."},
    "approval_request": {"subject": "Approval requested — Order {order_id}", "intro": "We are awaiting your confirmation to proceed with production."},
    "challan": {"subject": "Delivery challan — Order {order_id}", "intro": "Please find attached the dispatch challan for your records."},
    "delivery_pod": {"subject": "Proof of delivery — Order {order_id}", "intro": "Delivery has been completed. The proof of delivery is attached."},
    "invoice": {"subject": "Invoice {invoice_number} — Order {order_id}", "intro": "Please find your invoice attached. Kindly process at your earliest convenience."},
    "payment_reminder": {"subject": "Payment reminder — Invoice {invoice_number}", "intro": "This is a friendly reminder for the outstanding balance on the invoice below."},
    "generic": {"subject": "Update on Order {order_id}", "intro": "Please see the update below."},
}

# ---------- helpers ----------
def now(): return datetime.now(timezone.utc).isoformat()
def today(): return datetime.now(timezone.utc).strftime("%Y-%m-%d")
def clean(doc):
    if not doc: return None
    doc = dict(doc); doc.pop("_id", None); return doc
def hash_password(p): return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()
def verify_password(p, h):
    try: return bcrypt.checkpw(p.encode(), h.encode())
    except Exception: return False
def secret(): return os.environ["JWT_SECRET"]
def token(user_id, kind="access"):
    age = 3600 if kind == "access" else 604800
    return jwt.encode({"sub": user_id, "type": kind, "exp": datetime.now(timezone.utc) + timedelta(seconds=age)}, secret(), algorithm=JWT_ALGORITHM)

def new_id(prefix): return f"{prefix}-{secrets.token_hex(3).upper()}"

async def current_user(request: Request):
    value = request.cookies.get("access_token")
    auth = request.headers.get("Authorization", "")
    if not value and auth.startswith("Bearer "): value = auth[7:]
    if not value: raise HTTPException(401, "Please sign in to continue")
    try: payload = jwt.decode(value, secret(), algorithms=[JWT_ALGORITHM])
    except jwt.PyJWTError: raise HTTPException(401, "Your session has expired")
    user = await db.users.find_one({"user_id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user: raise HTTPException(401, "User not found")
    return user

def require(*roles):
    async def dep(user=Depends(current_user)):
        if roles and user["role"] not in roles and user["role"] != "admin":
            raise HTTPException(403, "You do not have access to this area")
        return user
    return dep

async def notify(roles, order_id, message):
    for r in roles:
        await db.notifications.insert_one({"notif_id": new_id("NT"), "role": r, "order_id": order_id, "message": message, "read": False, "created_at": now()})

async def add_history(order_id, event, by, remarks=""):
    await db.orders.update_one({"order_id": order_id}, {"$push": {"history": {"event": event, "by": by, "at": now(), "remarks": remarks}}, "$set": {"updated_at": now()}})

# ---------- models ----------
class Login(BaseModel):
    email: EmailStr
    password: str

class ClientCreate(BaseModel):
    company_name: str
    contact_person: str
    email: EmailStr
    phone: str
    alternate_phone: str = ""
    gstin: str = ""
    pan: str = ""
    billing_address: str = ""
    shipping_address: str = ""
    payment_terms: str = "Net 30"
    credit_limit: float = 0.0
    account_manager: str = ""
    status: str = "Active"
    notes: str = ""

class ContactCreate(BaseModel):
    client_id: str
    name: str
    designation: str = ""
    email: EmailStr
    phone: str = ""

class OrderCreate(BaseModel):
    client_id: str
    product: str
    product_code: str = ""
    quantity: int = Field(gt=0)
    unit_price: float = Field(gt=0)
    required_delivery_date: str
    po_number: str = ""
    po_date: str = ""
    priority: str = "Normal"
    notes: str = ""

class DesignCreate(BaseModel):
    designer: str = ""
    comments: str = ""
    file_id: str = ""
    file_name: str = ""

class ApprovalCreate(BaseModel):
    design_id: str = ""
    status: str  # Approved / Rejected / Revision Required
    approver: str
    comments: str = ""

class ProductionUpdate(BaseModel):
    status: str  # Not Started / In Progress / Completed / On Hold / Cancelled
    assigned_to: str = ""
    start_date: str = ""
    expected_completion: str = ""
    actual_completion: str = ""
    remarks: str = ""

class DispatchCreate(BaseModel):
    dispatch_date: str
    quantity: int = Field(gt=0)
    transporter: str
    tracking_number: str = ""
    delivery_address: str = ""
    remarks: str = ""

class ChallanCreate(BaseModel):
    challan_date: str
    quantity: int = Field(gt=0)
    transporter: str = ""
    tracking_number: str = ""
    file_id: str = ""

class DeliveryCreate(BaseModel):
    challan_id: str = ""
    expected_date: str = ""
    actual_date: str
    received_by: str
    pod_file_id: str = ""
    remarks: str = ""
    status: str = "Delivered"

class InvoiceCreate(BaseModel):
    invoice_date: str
    payment_terms: str = "Net 30"
    due_date: str = ""

class PaymentCreate(BaseModel):
    invoice_id: str
    amount: float = Field(gt=0)
    date: str
    mode: str = "Bank Transfer"
    reference: str = ""

# ---------- root ----------
@api.get("/")
async def root(): return {"message": "Pinaki Solutions CRM API"}

# ---------- auth ----------
@api.post("/auth/login")
async def login(data: Login, response: Response):
    email = data.email.lower()
    attempt = await db.login_attempts.find_one({"email": email}, {"_id": 0})
    if attempt and attempt.get("locked_until", 0) > datetime.now(timezone.utc).timestamp():
        raise HTTPException(429, "Too many failed attempts. Try again in 15 minutes")
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        failures = (attempt or {}).get("failures", 0) + 1
        locked = datetime.now(timezone.utc).timestamp() + 900 if failures >= 5 else 0
        await db.login_attempts.update_one({"email": email}, {"$set": {"email": email, "failures": failures, "locked_until": locked}}, upsert=True)
        raise HTTPException(401, "Incorrect email or password")
    await db.login_attempts.delete_one({"email": email})
    response.set_cookie("access_token", token(user["user_id"]), httponly=True, secure=True, samesite="none", max_age=3600)
    response.set_cookie("refresh_token", token(user["user_id"], "refresh"), httponly=True, secure=True, samesite="none", max_age=604800)
    return clean({k: v for k, v in user.items() if k != "password_hash"})

@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token")
    response.delete_cookie("refresh_token")
    return {"ok": True}

@api.get("/auth/me")
async def me(user=Depends(current_user)): return user

# ---------- dashboard ----------
@api.get("/dashboard")
async def dashboard(user=Depends(current_user)):
    orders = await db.orders.find({}, {"_id": 0}).to_list(1000)
    invoices = await db.invoices.find({}, {"_id": 0}).to_list(1000)
    payments = await db.payments.find({}, {"_id": 0}).to_list(1000)
    stage_counts = {stage: sum(1 for o in orders if o.get("current_stage") == stage) for stage in STAGES}
    outstanding = sum(i.get("total", 0) - sum(p.get("amount", 0) for p in payments if p.get("invoice_id") == i.get("invoice_id")) for i in invoices)
    overdue_amt = 0
    today_ts = datetime.now(timezone.utc).timestamp()
    for i in invoices:
        paid = sum(p.get("amount", 0) for p in payments if p.get("invoice_id") == i.get("invoice_id"))
        if paid < i.get("total", 0):
            try:
                due = datetime.fromisoformat(i.get("due_date", "1970-01-01")).timestamp()
                if due < today_ts: overdue_amt += i.get("total", 0) - paid
            except Exception: pass
    # monthly aggregation from orders
    months = {}
    for o in orders:
        m = (o.get("order_date") or "")[:7] or "unknown"
        months.setdefault(m, {"month": m, "orders": 0, "revenue": 0})
        months[m]["orders"] += 1
        months[m]["revenue"] += o.get("total_value", 0)
    monthly = sorted(months.values(), key=lambda x: x["month"])[-6:]
    delivery_counts = {}
    for s in ["Pending", "In Transit", "Delivered", "Failed"]:
        delivery_counts[s] = sum(1 for o in orders if o.get("delivery_status") == s)
    return {
        "total_orders": len(orders),
        "stage_counts": stage_counts,
        "new_orders": stage_counts.get("Order Received", 0),
        "awaiting_design": stage_counts.get("Design", 0),
        "awaiting_approval": stage_counts.get("Client Approval", 0),
        "awaiting_production": sum(1 for o in orders if o.get("approval_status") == "Approved" and o.get("production_status") == "Not Started"),
        "in_production": sum(1 for o in orders if o.get("production_status") == "In Progress"),
        "ready_dispatch": sum(1 for o in orders if o.get("production_status") == "Completed" and o.get("dispatch_status") == "Pending"),
        "dispatched": sum(1 for o in orders if o.get("dispatch_status") == "Dispatched"),
        "awaiting_delivery": sum(1 for o in orders if o.get("dispatch_status") == "Dispatched" and o.get("delivery_status") != "Delivered"),
        "delivered": sum(1 for o in orders if o.get("delivery_status") == "Delivered"),
        "pending_invoices": sum(1 for o in orders if o.get("delivery_status") == "Delivered" and o.get("invoice_status") in ["Pending", None]),
        "generated_invoices": len(invoices),
        "outstanding": round(outstanding, 2),
        "overdue": round(overdue_amt, 2),
        "monthly": monthly,
        "delivery_counts": delivery_counts,
    }

# ---------- clients ----------
@api.get("/clients")
async def list_clients(user=Depends(current_user)):
    return await db.clients.find({}, {"_id": 0}).sort("company_name", 1).to_list(1000)

@api.post("/clients")
async def create_client(data: ClientCreate, user=Depends(require("admin", "sales"))):
    doc = data.model_dump()
    doc.update({"client_id": new_id("CL"), "created_at": now()})
    await db.clients.insert_one(doc)
    return clean(doc)

@api.get("/clients/{client_id}")
async def client_detail(client_id: str, user=Depends(current_user)):
    c = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    if not c: raise HTTPException(404, "Client not found")
    c["orders"] = await db.orders.find({"client_id": client_id}, {"_id": 0}).sort("order_date", -1).to_list(500)
    c["contacts"] = await db.contacts.find({"client_id": client_id}, {"_id": 0}).to_list(500)
    c["invoices"] = await db.invoices.find({"client_id": client_id}, {"_id": 0}).to_list(500)
    return c

# ---------- contacts ----------
@api.get("/contacts")
async def list_contacts(client_id: str = "", user=Depends(current_user)):
    q = {"client_id": client_id} if client_id else {}
    return await db.contacts.find(q, {"_id": 0}).to_list(1000)

@api.post("/contacts")
async def create_contact(data: ContactCreate, user=Depends(require("admin", "sales"))):
    if not await db.clients.find_one({"client_id": data.client_id}, {"_id": 1}):
        raise HTTPException(404, "Client not found")
    doc = data.model_dump()
    doc.update({"contact_id": new_id("CT"), "created_at": now()})
    await db.contacts.insert_one(doc)
    return clean(doc)

# ---------- orders ----------
@api.get("/orders")
async def list_orders(search: str = "", stage: str = "", client_id: str = "", user=Depends(current_user)):
    q = {}
    if stage: q["current_stage"] = stage
    if client_id: q["client_id"] = client_id
    if search:
        q["$or"] = [
            {"order_id": {"$regex": search, "$options": "i"}},
            {"client_name": {"$regex": search, "$options": "i"}},
            {"product": {"$regex": search, "$options": "i"}},
            {"po_number": {"$regex": search, "$options": "i"}},
        ]
    return await db.orders.find(q, {"_id": 0}).sort("order_date", -1).to_list(1000)

@api.post("/orders")
async def create_order(data: OrderCreate, user=Depends(require("admin", "sales"))):
    c = await db.clients.find_one({"client_id": data.client_id}, {"_id": 0})
    if not c: raise HTTPException(404, "Client not found")
    subtotal = data.quantity * data.unit_price
    gst = round(subtotal * 0.18, 2)
    order_id = f"PS-{datetime.now().strftime('%y%m')}-{secrets.token_hex(2).upper()}"
    doc = data.model_dump()
    doc.update({
        "order_id": order_id,
        "client_id": data.client_id,
        "client_name": c["company_name"],
        "order_date": today(),
        "order_value": subtotal,
        "gst": gst,
        "total_value": subtotal + gst,
        "current_stage": "Order Received",
        "approval_status": "Pending Approval",
        "production_status": "Not Started",
        "dispatch_status": "Pending",
        "delivery_status": "Pending",
        "invoice_status": "Pending",
        "outstanding_amount": subtotal + gst,
        "sales_owner": user["name"],
        "history": [{"event": "Order created", "by": user["name"], "at": now(), "remarks": ""}],
        "created_at": now(),
    })
    await db.orders.insert_one(doc)
    await notify(["design", "sales"], order_id, f"New order {order_id} received from {c['company_name']}")
    return clean(doc)

@api.get("/orders/{order_id}")
async def order_detail(order_id: str, user=Depends(current_user)):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    o["designs"] = await db.designs.find({"order_id": order_id}, {"_id": 0}).sort("version", -1).to_list(100)
    o["approvals"] = await db.approvals.find({"order_id": order_id}, {"_id": 0}).sort("created_at", -1).to_list(100)
    o["dispatches"] = await db.dispatches.find({"order_id": order_id}, {"_id": 0}).to_list(100)
    o["challans"] = await db.challans.find({"order_id": order_id}, {"_id": 0}).to_list(100)
    o["deliveries"] = await db.deliveries.find({"order_id": order_id}, {"_id": 0}).to_list(100)
    o["invoices"] = await db.invoices.find({"order_id": order_id}, {"_id": 0}).to_list(100)
    o["payments"] = await db.payments.find({"order_id": order_id}, {"_id": 0}).to_list(100)
    o["documents"] = await db.documents.find({"order_id": order_id}, {"_id": 0}).sort("created_at", -1).to_list(200)
    return o

class StageUpdate(BaseModel):
    stage: str
    status: Optional[str] = None
    remarks: str = ""

@api.patch("/orders/{order_id}/stage")
async def update_stage(order_id: str, data: StageUpdate, user=Depends(current_user)):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    if data.stage == "Production" and o.get("approval_status") != "Approved":
        raise HTTPException(409, "Production is locked until client approval is Approved")
    if data.stage in ["Dispatch", "Challan"] and o.get("production_status") != "Completed":
        raise HTTPException(409, "Dispatch is locked until production is Completed")
    if data.stage in ["Invoice", "Accounts"] and o.get("delivery_status") != "Delivered":
        raise HTTPException(409, "Invoice is locked until delivery is Delivered")
    await db.orders.update_one({"order_id": order_id}, {"$set": {"current_stage": data.stage, "updated_at": now()}})
    await add_history(order_id, f"Moved to {data.stage}", user["name"], data.remarks)
    return await db.orders.find_one({"order_id": order_id}, {"_id": 0})

# ---------- designs ----------
@api.get("/designs")
async def list_designs(user=Depends(current_user)):
    return await db.designs.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)

@api.post("/orders/{order_id}/designs")
async def create_design(order_id: str, data: DesignCreate, user=Depends(require("admin", "sales", "design"))):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    version = await db.designs.count_documents({"order_id": order_id}) + 1
    doc = data.model_dump()
    doc.update({
        "design_id": new_id("DS"),
        "order_id": order_id,
        "client_id": o["client_id"],
        "version": version,
        "designer": data.designer or user["name"],
        "status": "Shared",
        "created_at": now(),
    })
    await db.designs.insert_one(doc)
    updates = {"current_stage": "Client Approval", "approval_status": "Pending Approval"}
    await db.orders.update_one({"order_id": order_id}, {"$set": updates})
    await add_history(order_id, f"Design V{version} uploaded", user["name"], data.comments)
    await notify(["sales"], order_id, f"Design V{version} ready for {o['client_name']} approval")
    return clean(doc)

# ---------- approvals ----------
@api.get("/approvals")
async def list_approvals(user=Depends(current_user)):
    return await db.approvals.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)

@api.post("/orders/{order_id}/approvals")
async def create_approval(order_id: str, data: ApprovalCreate, user=Depends(require("admin", "sales", "design"))):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    if data.status not in ["Approved", "Rejected", "Revision Required"]:
        raise HTTPException(400, "Approval status must be Approved, Rejected or Revision Required")
    doc = data.model_dump()
    doc.update({"approval_id": new_id("AP"), "order_id": order_id, "created_at": now()})
    await db.approvals.insert_one(doc)
    updates = {"approval_status": data.status, "updated_at": now()}
    if data.status == "Approved":
        updates["current_stage"] = "Production"
        await notify(["production"], order_id, f"{order_id} approved — production can start")
    else:
        updates["current_stage"] = "Design"
        await notify(["design"], order_id, f"{order_id} needs {data.status.lower()}")
    await db.orders.update_one({"order_id": order_id}, {"$set": updates})
    await add_history(order_id, f"Client approval: {data.status}", data.approver or user["name"], data.comments)
    return clean(doc)

# ---------- production ----------
@api.patch("/orders/{order_id}/production")
async def update_production(order_id: str, data: ProductionUpdate, user=Depends(require("admin", "production"))):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    if o.get("approval_status") != "Approved":
        raise HTTPException(409, "Production is locked until client approval is Approved")
    if data.status not in ["Not Started", "In Progress", "Completed", "On Hold", "Cancelled"]:
        raise HTTPException(400, "Invalid production status")
    updates = {"production_status": data.status, "updated_at": now()}
    if data.assigned_to: updates["assigned_producer"] = data.assigned_to
    if data.start_date: updates["production_start_date"] = data.start_date
    if data.expected_completion: updates["production_expected"] = data.expected_completion
    if data.actual_completion: updates["production_actual"] = data.actual_completion
    if data.status == "In Progress": updates["current_stage"] = "Production"
    if data.status == "Completed":
        updates["current_stage"] = "Dispatch"
        updates["production_actual"] = data.actual_completion or today()
        await notify(["dispatch"], order_id, f"{order_id} production completed — ready for dispatch")
    await db.orders.update_one({"order_id": order_id}, {"$set": updates})
    await add_history(order_id, f"Production: {data.status}", user["name"], data.remarks)
    return await db.orders.find_one({"order_id": order_id}, {"_id": 0})

# ---------- dispatch ----------
@api.get("/dispatches")
async def list_dispatches(user=Depends(current_user)):
    return await db.dispatches.find({}, {"_id": 0}).sort("dispatch_date", -1).to_list(500)

@api.post("/orders/{order_id}/dispatches")
async def create_dispatch(order_id: str, data: DispatchCreate, user=Depends(require("admin", "dispatch"))):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    if o.get("production_status") != "Completed":
        raise HTTPException(409, "Dispatch is locked until production is Completed")
    doc = data.model_dump()
    doc.update({"dispatch_id": new_id("DP"), "order_id": order_id, "client_name": o["client_name"], "status": "Dispatched", "created_at": now()})
    await db.dispatches.insert_one(doc)
    # auto-create challan
    challan_number = f"CH-{datetime.now().strftime('%y%m')}-{secrets.token_hex(2).upper()}"
    challan = {
        "challan_id": new_id("CN"), "order_id": order_id, "client_name": o["client_name"],
        "challan_number": challan_number, "challan_date": data.dispatch_date,
        "quantity": data.quantity, "transporter": data.transporter,
        "tracking_number": data.tracking_number, "created_at": now(),
    }
    await db.challans.insert_one(challan)
    # auto-generate challan PDF and attach as document
    try:
        client = await db.clients.find_one({"client_id": o["client_id"]}, {"_id": 0})
        pdf_bytes = generate_challan_pdf(o, challan, client)
        pdf_name = f"{challan_number}.pdf"
        gridfs_id = await gridfs.upload_from_stream(f"{order_id}-{secrets.token_hex(3)}-{pdf_name}", pdf_bytes, metadata={"order_id": order_id, "content_type": "application/pdf"})
        doc_record = {
            "document_id": new_id("DOC"),
            "order_id": order_id,
            "name": pdf_name,
            "category": "Challan",
            "content_type": "application/pdf",
            "storage_id": str(gridfs_id),
            "storage_backend": "mongo-gridfs",
            "uploaded_by": user["name"],
            "created_at": now(),
            "challan_id": challan["challan_id"],
        }
        await db.documents.insert_one(doc_record)
        await db.challans.update_one({"challan_id": challan["challan_id"]}, {"$set": {"pdf_document_id": doc_record["document_id"]}})
    except Exception:
        logging.exception("Failed to generate challan PDF")
    await db.orders.update_one({"order_id": order_id}, {"$set": {
        "current_stage": "Delivery", "dispatch_status": "Dispatched",
        "delivery_status": "In Transit", "dispatch_date": data.dispatch_date,
        "challan_number": challan_number, "tracking_number": data.tracking_number,
    }})
    await add_history(order_id, f"Dispatched via {data.transporter}", user["name"], data.remarks)
    await notify(["accounts", "sales"], order_id, f"{order_id} dispatched — challan {challan_number}")
    return clean(doc)

# ---------- challans ----------
@api.get("/challans")
async def list_challans(user=Depends(current_user)):
    return await db.challans.find({}, {"_id": 0}).sort("challan_date", -1).to_list(500)

# ---------- delivery ----------
@api.get("/deliveries")
async def list_deliveries(user=Depends(current_user)):
    return await db.deliveries.find({}, {"_id": 0}).sort("actual_date", -1).to_list(500)

@api.post("/orders/{order_id}/deliveries")
async def create_delivery(order_id: str, data: DeliveryCreate, user=Depends(require("admin", "dispatch", "accounts"))):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    if o.get("dispatch_status") != "Dispatched":
        raise HTTPException(409, "Delivery requires dispatch to be completed")
    doc = data.model_dump()
    doc.update({"delivery_id": new_id("DL"), "order_id": order_id, "client_name": o["client_name"], "created_at": now()})
    await db.deliveries.insert_one(doc)
    updates = {"delivery_status": data.status, "delivery_date": data.actual_date, "updated_at": now()}
    if data.status == "Delivered":
        updates["current_stage"] = "Accounts"
        await notify(["accounts"], order_id, f"{order_id} delivered — ready for invoicing")
    await db.orders.update_one({"order_id": order_id}, {"$set": updates})
    await add_history(order_id, f"Delivery: {data.status}", user["name"], data.remarks)
    return clean(doc)

# ---------- invoices ----------
@api.get("/invoices")
async def list_invoices(user=Depends(current_user)):
    invoices = await db.invoices.find({}, {"_id": 0}).sort("invoice_date", -1).to_list(500)
    payments = await db.payments.find({}, {"_id": 0}).to_list(1000)
    for inv in invoices:
        paid = sum(p.get("amount", 0) for p in payments if p.get("invoice_id") == inv.get("invoice_id"))
        inv["amount_received"] = paid
        inv["outstanding"] = round(inv.get("total", 0) - paid, 2)
        if paid >= inv.get("total", 0):
            inv["status"] = "Paid"
        elif paid > 0:
            inv["status"] = "Partially Paid"
        else:
            try:
                due = datetime.fromisoformat(inv.get("due_date", "1970-01-01")).timestamp()
                inv["status"] = "Overdue" if due < datetime.now(timezone.utc).timestamp() else "Due"
            except Exception:
                inv["status"] = inv.get("status", "Due")
    return invoices

@api.post("/orders/{order_id}/invoices")
async def create_invoice(order_id: str, data: InvoiceCreate, user=Depends(require("admin", "accounts"))):
    o = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not o: raise HTTPException(404, "Order not found")
    if o.get("delivery_status") != "Delivered":
        raise HTTPException(409, "Invoice is locked until delivery is Delivered")
    if await db.invoices.find_one({"order_id": order_id}, {"_id": 1}):
        raise HTTPException(409, "Invoice already generated for this order")
    invoice_number = f"INV-{datetime.now().strftime('%y%m')}-{secrets.token_hex(2).upper()}"
    doc = {
        "invoice_id": new_id("IV"),
        "invoice_number": invoice_number,
        "order_id": order_id,
        "client_id": o["client_id"],
        "client_name": o["client_name"],
        "invoice_date": data.invoice_date,
        "due_date": data.due_date or data.invoice_date,
        "payment_terms": data.payment_terms,
        "taxable": o.get("order_value", 0),
        "gst": o.get("gst", 0),
        "total": o.get("total_value", 0),
        "status": "Generated",
        "created_at": now(),
    }
    await db.invoices.insert_one(doc)
    # auto-generate PDF and store in GridFS
    try:
        client = await db.clients.find_one({"client_id": o["client_id"]}, {"_id": 0})
        pdf_bytes = generate_invoice_pdf(o, doc, client)
        pdf_name = f"{invoice_number}.pdf"
        gridfs_id = await gridfs.upload_from_stream(f"{order_id}-{secrets.token_hex(3)}-{pdf_name}", pdf_bytes, metadata={"order_id": order_id, "content_type": "application/pdf"})
        doc_record = {
            "document_id": new_id("DOC"),
            "order_id": order_id,
            "name": pdf_name,
            "category": "Invoice",
            "content_type": "application/pdf",
            "storage_id": str(gridfs_id),
            "storage_backend": "mongo-gridfs",
            "uploaded_by": user["name"],
            "created_at": now(),
            "invoice_id": doc["invoice_id"],
        }
        await db.documents.insert_one(doc_record)
        await db.invoices.update_one({"invoice_id": doc["invoice_id"]}, {"$set": {"pdf_document_id": doc_record["document_id"]}})
    except Exception:
        logging.exception("Failed to generate invoice PDF")
    await db.orders.update_one({"order_id": order_id}, {"$set": {"invoice_status": "Generated", "invoice_number": invoice_number, "current_stage": "Payment"}})
    await add_history(order_id, f"Invoice {invoice_number} generated with PDF", user["name"])
    await notify(["sales", "accounts"], order_id, f"Invoice {invoice_number} generated for {o['client_name']}")
    return clean(doc)

# ---------- payments ----------
@api.get("/payments")
async def list_payments(user=Depends(current_user)):
    return await db.payments.find({}, {"_id": 0}).sort("date", -1).to_list(500)

@api.post("/payments")
async def create_payment(data: PaymentCreate, user=Depends(require("admin", "accounts"))):
    inv = await db.invoices.find_one({"invoice_id": data.invoice_id}, {"_id": 0})
    if not inv: raise HTTPException(404, "Invoice not found")
    doc = data.model_dump()
    doc.update({
        "payment_id": new_id("PY"),
        "order_id": inv["order_id"],
        "client_id": inv["client_id"],
        "client_name": inv["client_name"],
        "invoice_number": inv["invoice_number"],
        "created_at": now(),
    })
    await db.payments.insert_one(doc)
    # update outstanding on order
    payments = await db.payments.find({"invoice_id": data.invoice_id}, {"_id": 0}).to_list(500)
    paid = sum(p["amount"] for p in payments)
    outstanding = max(0, inv["total"] - paid)
    status = "Paid" if outstanding == 0 else "Partially Paid"
    await db.invoices.update_one({"invoice_id": data.invoice_id}, {"$set": {"status": status}})
    await db.orders.update_one({"order_id": inv["order_id"]}, {"$set": {"outstanding_amount": outstanding, "invoice_status": status}})
    await add_history(inv["order_id"], f"Payment ₹{data.amount} received via {data.mode}", user["name"], data.reference)
    return clean(doc)

# ---------- documents ----------
@api.post("/orders/{order_id}/documents")
async def upload_document(order_id: str, file: UploadFile = File(...), category: str = "General", user=Depends(current_user)):
    if not await db.orders.find_one({"order_id": order_id}, {"_id": 1}):
        raise HTTPException(404, "Order not found")
    safe = f"{order_id}-{secrets.token_hex(4)}-{file.filename}"
    file_id = await gridfs.upload_from_stream(safe, await file.read(), metadata={"order_id": order_id, "content_type": file.content_type or "application/octet-stream"})
    doc = {
        "document_id": new_id("DOC"),
        "order_id": order_id,
        "name": file.filename,
        "category": category,
        "content_type": file.content_type or "application/octet-stream",
        "storage_id": str(file_id),
        "storage_backend": "mongo-gridfs",
        "uploaded_by": user["name"],
        "created_at": now(),
    }
    await db.documents.insert_one(doc)
    await add_history(order_id, f"Document '{file.filename}' uploaded ({category})", user["name"])
    return clean(doc)

@api.get("/orders/{order_id}/documents")
async def list_order_documents(order_id: str, user=Depends(current_user)):
    return await db.documents.find({"order_id": order_id}, {"_id": 0}).sort("created_at", -1).to_list(200)

@api.get("/documents/{document_id}/download")
async def download_document(document_id: str, user=Depends(current_user)):
    d = await db.documents.find_one({"document_id": document_id}, {"_id": 0})
    if not d: raise HTTPException(404, "Document not found")
    try:
        stream = await gridfs.open_download_stream(ObjectId(d["storage_id"]))
    except Exception:
        raise HTTPException(404, "File not found in storage")
    data = await stream.read()

    async def gen():
        yield data
    return StreamingResponse(gen(), media_type=d.get("content_type", "application/octet-stream"), headers={"Content-Disposition": f'attachment; filename="{d["name"]}"'})

# ---------- notifications ----------
@api.get("/notifications")
async def list_notifications(user=Depends(current_user)):
    role = user["role"]
    q = {} if role == "admin" else {"role": role}
    return await db.notifications.find(q, {"_id": 0}).sort("created_at", -1).limit(50).to_list(50)

@api.post("/notifications/{notif_id}/read")
async def mark_read(notif_id: str, user=Depends(current_user)):
    await db.notifications.update_one({"notif_id": notif_id}, {"$set": {"read": True}})
    return {"ok": True}

# ---------- search ----------
@api.get("/search")
async def global_search(q: str, user=Depends(current_user)):
    if not q: return {"orders": [], "clients": [], "invoices": [], "challans": []}
    rx = {"$regex": q, "$options": "i"}
    return {
        "orders": await db.orders.find({"$or": [{"order_id": rx}, {"client_name": rx}, {"product": rx}, {"po_number": rx}]}, {"_id": 0}).limit(20).to_list(20),
        "clients": await db.clients.find({"$or": [{"company_name": rx}, {"contact_person": rx}, {"email": rx}]}, {"_id": 0}).limit(20).to_list(20),
        "invoices": await db.invoices.find({"$or": [{"invoice_number": rx}, {"client_name": rx}]}, {"_id": 0}).limit(20).to_list(20),
        "challans": await db.challans.find({"$or": [{"challan_number": rx}, {"tracking_number": rx}]}, {"_id": 0}).limit(20).to_list(20),
    }

# ---------- reports ----------
@api.get("/reports")
async def reports(user=Depends(current_user)):
    orders = await db.orders.find({}, {"_id": 0}).to_list(1000)
    invoices = await db.invoices.find({}, {"_id": 0}).to_list(1000)
    payments = await db.payments.find({}, {"_id": 0}).to_list(1000)
    by_client = {}
    by_product = {}
    for o in orders:
        by_client.setdefault(o["client_name"], {"client": o["client_name"], "orders": 0, "revenue": 0})
        by_client[o["client_name"]]["orders"] += 1
        by_client[o["client_name"]]["revenue"] += o.get("total_value", 0)
        by_product.setdefault(o["product"], {"product": o["product"], "orders": 0, "revenue": 0})
        by_product[o["product"]]["orders"] += 1
        by_product[o["product"]]["revenue"] += o.get("total_value", 0)
    def delta_days(a, b):
        try: return (datetime.fromisoformat(b) - datetime.fromisoformat(a)).days
        except Exception: return None
    otd, oti = [], []
    for o in orders:
        d1 = delta_days(o.get("order_date", ""), o.get("delivery_date", ""))
        d2 = delta_days(o.get("order_date", ""), (o.get("invoice_number") and next((i["invoice_date"] for i in invoices if i["order_id"] == o["order_id"]), "")) or "")
        if d1 is not None: otd.append(d1)
        if d2 is not None: oti.append(d2)
    return {
        "order_volume": len(orders),
        "revenue": sum(o.get("total_value", 0) for o in orders),
        "by_client": sorted(by_client.values(), key=lambda x: -x["revenue"])[:10],
        "by_product": sorted(by_product.values(), key=lambda x: -x["revenue"])[:10],
        "pending_approvals": sum(1 for o in orders if o.get("approval_status") == "Pending Approval"),
        "production_pending": sum(1 for o in orders if o.get("approval_status") == "Approved" and o.get("production_status") != "Completed"),
        "dispatch_pending": sum(1 for o in orders if o.get("production_status") == "Completed" and o.get("dispatch_status") == "Pending"),
        "delivery_pending": sum(1 for o in orders if o.get("dispatch_status") == "Dispatched" and o.get("delivery_status") != "Delivered"),
        "invoice_pending": sum(1 for o in orders if o.get("delivery_status") == "Delivered" and o.get("invoice_status") in ["Pending", None]),
        "outstanding_total": sum(i.get("total", 0) - sum(p["amount"] for p in payments if p["invoice_id"] == i["invoice_id"]) for i in invoices),
        "avg_order_to_delivery": round(sum(otd) / len(otd), 1) if otd else 0,
        "avg_order_to_invoice": round(sum(oti) / len(oti), 1) if oti else 0,
    }

# ---------- leads / sales pipeline ----------
class LeadCreate(BaseModel):
    company_name: str
    contact_person: str
    email: EmailStr
    phone: str = ""
    source: str = "Website"
    estimated_value: float = 0
    product_interest: str = ""
    assigned_to: str = ""
    next_follow_up: str = ""
    notes: str = ""

class LeadUpdate(BaseModel):
    status: Optional[str] = None
    company_name: Optional[str] = None
    contact_person: Optional[str] = None
    email: Optional[EmailStr] = None
    phone: Optional[str] = None
    source: Optional[str] = None
    estimated_value: Optional[float] = None
    product_interest: Optional[str] = None
    assigned_to: Optional[str] = None
    next_follow_up: Optional[str] = None
    notes: Optional[str] = None

class LeadActivity(BaseModel):
    type: str
    summary: str
    outcome: str = ""

class LeadConvert(BaseModel):
    gstin: str = ""
    pan: str = ""
    payment_terms: str = "Net 30"
    credit_limit: float = 0
    billing_address: str = ""
    shipping_address: str = ""
    create_order: bool = False
    order_product: str = ""
    order_quantity: int = 0
    order_unit_price: float = 0
    order_required_delivery: str = ""

@api.get("/leads/pipeline")
async def leads_pipeline(user=Depends(current_user)):
    leads = await db.leads.find({}, {"_id": 0}).sort("created_at", -1).to_list(2000)
    buckets = {s: [] for s in LEAD_STATUSES}
    for l in leads:
        buckets.setdefault(l.get("status", "New"), []).append(l)
    totals = {s: round(sum(l.get("estimated_value", 0) for l in buckets[s]), 2) for s in LEAD_STATUSES}
    won_leads = buckets.get("Won", [])
    lost_leads = buckets.get("Lost", [])
    closed = len(won_leads) + len(lost_leads)
    win_rate = round(100 * len(won_leads) / closed, 1) if closed else 0.0
    open_value = round(sum(v for s, v in totals.items() if s not in ("Won", "Lost")), 2)
    return {"buckets": buckets, "totals": totals, "win_rate": win_rate, "open_value": open_value}

LEAD_IMPORT_COLUMNS = ["company_name", "contact_person", "email", "phone", "source", "estimated_value", "product_interest", "assigned_to", "next_follow_up", "status", "notes"]
LEAD_COLUMN_ALIASES = {"company": "company_name", "contact": "contact_person", "phone_number": "phone", "value": "estimated_value", "product": "product_interest", "owner": "assigned_to", "follow_up": "next_follow_up", "follow-up": "next_follow_up", "next_followup": "next_follow_up"}

@api.get("/leads/import-template")
async def leads_import_template(user=Depends(current_user)):
    def gen():
        yield (",".join(LEAD_IMPORT_COLUMNS) + "\n").encode()
        yield ("Acme Corp,John Doe,john@acme.example,+91 9000011111,Website,150000,Custom boxes,Riya Shah,2026-03-01,New,Initial outreach\n").encode()
    return StreamingResponse(gen(), media_type="text/csv", headers={"Content-Disposition": 'attachment; filename="leads_template.csv"'})

@api.post("/leads/import")
async def import_leads(file: UploadFile = File(...), user=Depends(require("admin", "sales"))):
    filename = file.filename or "upload"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    content = await file.read()
    rows = []
    try:
        if ext == "csv" or file.content_type == "text/csv":
            text = content.decode("utf-8-sig", errors="ignore")
            reader = csv.DictReader(io.StringIO(text))
            rows = [dict(r) for r in reader]
        elif ext in ("xlsx", "xlsm"):
            try:
                from openpyxl import load_workbook
            except ImportError:
                raise HTTPException(400, "Excel support requires openpyxl on the server")
            wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            ws = wb.active
            it = ws.iter_rows(values_only=True)
            first = next(it, None)
            if not first:
                raise HTTPException(400, "The file appears to be empty")
            headers = [str(h).strip() if h is not None else "" for h in first]
            for r in it:
                if all(c is None or str(c).strip() == "" for c in r):
                    continue
                rows.append({headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers)) if headers[i]})
        else:
            raise HTTPException(400, "Please upload a .csv or .xlsx file")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Could not read the file: {str(e)[:120]}")

    inserted, skipped = 0, 0
    errors = []
    for idx, raw_row in enumerate(rows, start=2):
        normalized = {}
        for k, v in (raw_row or {}).items():
            if not k:
                continue
            key = str(k).strip().lower().replace(" ", "_").replace("-", "_")
            key = LEAD_COLUMN_ALIASES.get(key, key)
            if isinstance(v, str):
                v = v.strip()
            normalized[key] = v
        try:
            company = normalized.get("company_name") or ""
            contact = normalized.get("contact_person") or ""
            email = normalized.get("email") or ""
            if not company or not contact or not email:
                errors.append({"row": idx, "message": "Missing required company_name / contact_person / email"})
                skipped += 1
                continue
            status = normalized.get("status") or "New"
            if status not in LEAD_STATUSES:
                status = "New"
            try:
                value = float(normalized.get("estimated_value") or 0)
            except Exception:
                value = 0.0
            follow_up = normalized.get("next_follow_up") or ""
            if follow_up and hasattr(follow_up, "isoformat"):
                follow_up = follow_up.isoformat()[:10]
            follow_up = str(follow_up)[:10] if follow_up else ""
            doc = {
                "lead_id": new_id("LD"),
                "company_name": str(company),
                "contact_person": str(contact),
                "email": str(email),
                "phone": str(normalized.get("phone") or ""),
                "source": str(normalized.get("source") or "Import"),
                "estimated_value": value,
                "product_interest": str(normalized.get("product_interest") or ""),
                "assigned_to": str(normalized.get("assigned_to") or user["name"]),
                "next_follow_up": follow_up,
                "notes": str(normalized.get("notes") or ""),
                "status": status,
                "activities": [{"type": "note", "summary": f"Imported from {filename}", "by": user["name"], "at": now()}],
                "created_by": user["name"],
                "created_at": now(),
                "import_source": filename,
            }
            await db.leads.insert_one(doc)
            inserted += 1
        except Exception as e:
            errors.append({"row": idx, "message": str(e)[:160]})
            skipped += 1
    logging.info(f"Lead import: {inserted}/{len(rows)} inserted from {filename}")
    return {"filename": filename, "total": len(rows), "inserted": inserted, "skipped": skipped, "errors": errors[:25], "columns_expected": LEAD_IMPORT_COLUMNS}

@api.get("/leads")
async def list_leads(status: str = "", assigned: str = "", search: str = "", user=Depends(current_user)):
    q = {}
    if status: q["status"] = status
    if assigned: q["assigned_to"] = assigned
    if search:
        rx = {"$regex": search, "$options": "i"}
        q["$or"] = [{"company_name": rx}, {"contact_person": rx}, {"email": rx}, {"product_interest": rx}, {"lead_id": rx}]
    return await db.leads.find(q, {"_id": 0}).sort("created_at", -1).to_list(1000)

@api.post("/leads")
async def create_lead(data: LeadCreate, user=Depends(require("admin", "sales"))):
    doc = data.model_dump()
    doc.update({
        "lead_id": new_id("LD"),
        "status": "New",
        "activities": [{"type": "note", "summary": "Lead created", "by": user["name"], "at": now()}],
        "assigned_to": data.assigned_to or user["name"],
        "created_by": user["name"],
        "created_at": now(),
    })
    await db.leads.insert_one(doc)
    await notify(["sales"], None, f"New lead: {data.company_name}")
    return clean(doc)

@api.get("/leads/{lead_id}")
async def lead_detail(lead_id: str, user=Depends(current_user)):
    l = await db.leads.find_one({"lead_id": lead_id}, {"_id": 0})
    if not l: raise HTTPException(404, "Lead not found")
    if l.get("converted_client_id"):
        l["client"] = await db.clients.find_one({"client_id": l["converted_client_id"]}, {"_id": 0})
    return l

@api.patch("/leads/{lead_id}")
async def update_lead(lead_id: str, data: LeadUpdate, user=Depends(require("admin", "sales"))):
    updates = {k: v for k, v in data.model_dump().items() if v is not None}
    if not updates: raise HTTPException(400, "No changes provided")
    if "status" in updates and updates["status"] not in LEAD_STATUSES:
        raise HTTPException(400, "Invalid lead status")
    updates["updated_at"] = now()
    old = await db.leads.find_one({"lead_id": lead_id}, {"_id": 0})
    if not old: raise HTTPException(404, "Lead not found")
    push = None
    if "status" in updates and updates["status"] != old.get("status"):
        push = {"activities": {"type": "status_change", "summary": f"Moved from {old.get('status')} to {updates['status']}", "by": user["name"], "at": now()}}
    op = {"$set": updates}
    if push: op["$push"] = push
    await db.leads.update_one({"lead_id": lead_id}, op)
    return await db.leads.find_one({"lead_id": lead_id}, {"_id": 0})

@api.post("/leads/{lead_id}/activities")
async def add_lead_activity(lead_id: str, data: LeadActivity, user=Depends(require("admin", "sales"))):
    if data.type not in ("call", "email", "meeting", "note", "status_change"):
        raise HTTPException(400, "Invalid activity type")
    entry = {**data.model_dump(), "by": user["name"], "at": now()}
    r = await db.leads.update_one({"lead_id": lead_id}, {"$push": {"activities": entry}, "$set": {"updated_at": now()}})
    if r.matched_count == 0: raise HTTPException(404, "Lead not found")
    return entry

@api.post("/leads/{lead_id}/convert")
async def convert_lead(lead_id: str, data: LeadConvert, user=Depends(require("admin", "sales"))):
    l = await db.leads.find_one({"lead_id": lead_id}, {"_id": 0})
    if not l: raise HTTPException(404, "Lead not found")
    if l.get("converted_client_id"): raise HTTPException(409, "Lead already converted")
    client = {
        "client_id": new_id("CL"),
        "company_name": l["company_name"],
        "contact_person": l["contact_person"],
        "email": l["email"],
        "phone": l.get("phone", ""),
        "alternate_phone": "",
        "gstin": data.gstin, "pan": data.pan,
        "billing_address": data.billing_address,
        "shipping_address": data.shipping_address or data.billing_address,
        "payment_terms": data.payment_terms,
        "credit_limit": float(data.credit_limit or 0),
        "account_manager": l.get("assigned_to") or user["name"],
        "status": "Active",
        "notes": f"Converted from lead {l['lead_id']}. {l.get('notes','')}",
        "source_lead_id": l["lead_id"],
        "created_at": now(),
    }
    await db.clients.insert_one(client)
    await db.leads.update_one({"lead_id": lead_id}, {"$set": {
        "status": "Won", "converted_client_id": client["client_id"],
        "converted_at": now(), "converted_by": user["name"], "updated_at": now(),
    }, "$push": {"activities": {"type": "convert", "summary": f"Converted to client {client['company_name']}", "by": user["name"], "at": now()}}})
    await notify(["sales"], None, f"Lead {l['company_name']} converted to client")
    order = None
    if data.create_order and data.order_product and data.order_quantity > 0 and data.order_unit_price > 0 and data.order_required_delivery:
        subtotal = data.order_quantity * data.order_unit_price
        gst = round(subtotal * 0.18, 2)
        order = {
            "order_id": f"PS-{datetime.now().strftime('%y%m')}-{secrets.token_hex(2).upper()}",
            "client_id": client["client_id"], "client_name": client["company_name"],
            "product": data.order_product, "product_code": "",
            "quantity": data.order_quantity, "unit_price": data.order_unit_price,
            "order_date": today(),
            "order_value": subtotal, "gst": gst, "total_value": subtotal + gst,
            "required_delivery_date": data.order_required_delivery,
            "current_stage": "Order Received", "approval_status": "Pending Approval",
            "production_status": "Not Started", "dispatch_status": "Pending",
            "delivery_status": "Pending", "invoice_status": "Pending",
            "outstanding_amount": subtotal + gst,
            "sales_owner": l.get("assigned_to") or user["name"],
            "priority": "Normal",
            "source_lead_id": l["lead_id"],
            "history": [{"event": f"Order created from lead {l['lead_id']}", "by": user["name"], "at": now()}],
            "created_at": now(),
        }
        await db.orders.insert_one(order)
        await notify(["design", "sales"], order["order_id"], f"New order from converted lead {l['company_name']}")
    return {"client": clean(client), "order": clean(order) if order else None}

# ---------- emails ----------
class EmailSend(BaseModel):
    order_id: str
    recipient: EmailStr
    subject: str = ""
    message: str = ""
    document_ids: List[str] = []
    template: str = "generic"
    cc: List[EmailStr] = []

def build_email_html(template_key, order, message, sender_name, invoice=None, approval_url=None, pay_url=None):
    tpl = EMAIL_TEMPLATES.get(template_key, EMAIL_TEMPLATES["generic"])
    intro = tpl["intro"]
    rows = [
        ("Order ID", order.get("order_id")),
        ("Client", order.get("client_name")),
        ("Product", order.get("product")),
        ("Quantity", f"{order.get('quantity', 0):,} units"),
        ("Order value", f"₹{order.get('total_value', 0):,.0f}"),
        ("Required delivery", order.get("required_delivery_date", "—")),
    ]
    if invoice:
        rows += [
            ("Invoice #", invoice.get("invoice_number")),
            ("Invoice date", invoice.get("invoice_date")),
            ("Due date", invoice.get("due_date")),
            ("Total", f"₹{invoice.get('total', 0):,.0f}"),
        ]
    rows_html = "".join(f"<tr><td style='padding:8px 14px;color:#64748b;font-size:12px;'>{k}</td><td style='padding:8px 14px;color:#0f172a;font-size:13px;font-weight:600;'>{v}</td></tr>" for k, v in rows)
    body = (message or intro).replace("\n", "<br>")
    cta_html = ""
    if approval_url:
        cta_html += f"""
        <tr><td align='center' style='padding:8px 28px 8px;'>
          <a href='{approval_url}' style='display:inline-block;background:#2563eb;color:#fff;text-decoration:none;padding:14px 28px;border-radius:8px;font-weight:700;font-size:13px;letter-spacing:.3px;'>Review &amp; approve online</a>
          <div style='color:#94a3b8;font-size:11px;margin-top:8px;'>One-click approve, request revision or reject.</div>
        </td></tr>
        """
    if pay_url:
        cta_html += f"""
        <tr><td align='center' style='padding:8px 28px 24px;'>
          <a href='{pay_url}' style='display:inline-block;background:#16a34a;color:#fff;text-decoration:none;padding:14px 28px;border-radius:8px;font-weight:700;font-size:13px;letter-spacing:.3px;'>Pay online — view bank details</a>
          <div style='color:#94a3b8;font-size:11px;margin-top:8px;'>Secure link with bank transfer &amp; UPI details.</div>
        </td></tr>
        """
    return f"""
    <table width='100%' cellpadding='0' cellspacing='0' style='background:#f8fafc;padding:32px 0;font-family:Arial,sans-serif;'>
      <tr><td align='center'>
        <table width='560' cellpadding='0' cellspacing='0' style='background:#fff;border-radius:10px;overflow:hidden;box-shadow:0 4px 12px rgba(15,23,42,.08);'>
          <tr><td style='background:#0f172a;padding:24px 28px;'>
            <div style='color:#93c5fd;font-size:11px;letter-spacing:2px;font-weight:700;'>PINAKI SOLUTIONS</div>
            <div style='color:#fff;font-size:22px;font-weight:700;margin-top:4px;'>Order lifecycle update</div>
          </td></tr>
          <tr><td style='padding:26px 28px 8px;color:#0f172a;font-size:14px;line-height:1.6;'>{body}</td></tr>
          <tr><td style='padding:8px 28px 24px;'>
            <table cellpadding='0' cellspacing='0' style='width:100%;border-collapse:collapse;background:#f8fafc;border-radius:8px;overflow:hidden;'>
              {rows_html}
            </table>
          </td></tr>
          {cta_html}
          <tr><td style='padding:12px 28px 26px;color:#64748b;font-size:12px;line-height:1.6;border-top:1px solid #e2e8f0;'>
            Sent by <b style='color:#0f172a;'>{sender_name}</b> · Pinaki Solutions<br>
            <span style='color:#94a3b8;'>Reply to this email to reach the account team.</span>
          </td></tr>
        </table>
      </td></tr>
    </table>
    """

@api.post("/emails/send")
async def send_email(data: EmailSend, request: Request, user=Depends(current_user)):
    if not os.environ.get("RESEND_API_KEY"):
        raise HTTPException(400, "Email is not configured. Please add RESEND_API_KEY to the backend to enable client emails.")
    order = await db.orders.find_one({"order_id": data.order_id}, {"_id": 0})
    if not order:
        raise HTTPException(404, "Order not found")
    tpl = EMAIL_TEMPLATES.get(data.template, EMAIL_TEMPLATES["generic"])
    invoice = None
    if data.template in ("invoice", "payment_reminder"):
        invoice = await db.invoices.find_one({"order_id": data.order_id}, {"_id": 0})
    subject = data.subject or tpl["subject"].format(order_id=order.get("order_id"), invoice_number=(invoice or {}).get("invoice_number", ""))
    # auto-attach relevant PDFs
    doc_ids = list(data.document_ids)
    if data.template in ("invoice", "payment_reminder") and invoice and invoice.get("pdf_document_id") and invoice["pdf_document_id"] not in doc_ids:
        doc_ids.append(invoice["pdf_document_id"])
    if data.template == "challan":
        challan = await db.challans.find_one({"order_id": data.order_id}, {"_id": 0}, sort=[("created_at", -1)])
        if challan and challan.get("pdf_document_id") and challan["pdf_document_id"] not in doc_ids:
            doc_ids.append(challan["pdf_document_id"])
    attachments = []
    for doc_id in doc_ids:
        d = await db.documents.find_one({"document_id": doc_id}, {"_id": 0})
        if not d:
            continue
        try:
            stream = await gridfs.open_download_stream(ObjectId(d["storage_id"]))
            content = await stream.read()
            attachments.append({"filename": d["name"], "content": base64.b64encode(content).decode()})
        except Exception:
            continue
    # embed approval link for design_share
    approval_url = None
    if data.template == "design_share":
        latest_design = await db.designs.find_one({"order_id": data.order_id}, {"_id": 0}, sort=[("version", -1)])
        if latest_design:
            token = sign_approval_token(latest_design["design_id"], data.order_id)
            approval_url = f"{frontend_url(request)}/approve/{token}"
    # embed pay link for invoice/reminder
    pay_url = None
    if data.template in ("invoice", "payment_reminder") and invoice:
        pay_url = f"{frontend_url(request)}/pay/{sign_pay_token(invoice['invoice_id'])}"
    html = build_email_html(data.template, order, data.message, user["name"], invoice, approval_url, pay_url)
    params = {
        "from": f"Pinaki Solutions <{SENDER_EMAIL}>",
        "to": [data.recipient],
        "subject": subject,
        "html": html,
    }
    if data.cc:
        params["cc"] = data.cc
    if attachments:
        params["attachments"] = attachments
    try:
        result = await asyncio.to_thread(resend.Emails.send, params)
    except Exception as e:
        logging.exception("Resend send failed")
        raise HTTPException(502, f"Could not send email: {str(e)}")
    provider_id = result.get("id") if isinstance(result, dict) else getattr(result, "id", None)
    record = {
        "email_id": new_id("EM"),
        "order_id": data.order_id,
        "client_name": order.get("client_name"),
        "recipient": data.recipient,
        "cc": data.cc,
        "subject": subject,
        "template": data.template,
        "provider_id": provider_id,
        "document_ids": doc_ids,
        "sent_by": user["name"],
        "created_at": now(),
    }
    await db.emails.insert_one(record)
    await add_history(data.order_id, f"Email sent to {data.recipient} — {subject}", user["name"], data.template)
    return clean(record)

@api.get("/orders/{order_id}/emails")
async def order_emails(order_id: str, user=Depends(current_user)):
    return await db.emails.find({"order_id": order_id}, {"_id": 0}).sort("created_at", -1).to_list(200)

@api.get("/emails")
async def all_emails(user=Depends(current_user)):
    return await db.emails.find({}, {"_id": 0}).sort("created_at", -1).limit(200).to_list(200)

@api.get("/emails/config")
async def email_config(user=Depends(current_user)):
    return {"configured": bool(os.environ.get("RESEND_API_KEY")), "sender": SENDER_EMAIL}

# ---------- public payment page ----------
@api.get("/public/pay/{token}")
async def public_pay(token: str):
    try:
        payload = verify_pay_token(token)
    except jwt.PyJWTError:
        raise HTTPException(400, "This payment link is invalid or has expired")
    inv = await db.invoices.find_one({"invoice_id": payload["sub"]}, {"_id": 0})
    if not inv:
        raise HTTPException(404, "Invoice not found")
    order = await db.orders.find_one({"order_id": inv["order_id"]}, {"_id": 0})
    payments = await db.payments.find({"invoice_id": inv["invoice_id"]}, {"_id": 0}).to_list(100)
    paid = sum(p.get("amount", 0) for p in payments)
    outstanding = max(0, inv.get("total", 0) - paid)
    return {
        "invoice": {k: inv.get(k) for k in ("invoice_number", "invoice_date", "due_date", "payment_terms", "total", "status", "pdf_document_id")},
        "order": {k: (order or {}).get(k) for k in ("order_id", "client_name", "product", "quantity", "total_value", "po_number")},
        "amount_paid": paid,
        "outstanding": outstanding,
        "bank_details": BANK_DETAILS,
    }

# ---------- weekly management digest ----------
def ist_now():
    return datetime.now(timezone.utc) + timedelta(hours=5, minutes=30)

async def build_digest():
    orders = await db.orders.find({}, {"_id": 0}).to_list(2000)
    invoices = await db.invoices.find({}, {"_id": 0}).to_list(2000)
    payments = await db.payments.find({}, {"_id": 0}).to_list(3000)
    paid_by_invoice = {}
    for p in payments:
        paid_by_invoice[p["invoice_id"]] = paid_by_invoice.get(p["invoice_id"], 0) + p.get("amount", 0)
    outstanding = sum(max(0, i.get("total", 0) - paid_by_invoice.get(i["invoice_id"], 0)) for i in invoices)
    today = ist_now().date()
    overdue = 0
    for i in invoices:
        try:
            due = datetime.fromisoformat(i.get("due_date", "1970-01-01")).date()
        except Exception:
            continue
        if due < today and paid_by_invoice.get(i["invoice_id"], 0) < i.get("total", 0):
            overdue += i.get("total", 0) - paid_by_invoice.get(i["invoice_id"], 0)
    upcoming = []
    for o in orders:
        try:
            due = datetime.fromisoformat(o.get("required_delivery_date", "1970-01-01")).date()
        except Exception:
            continue
        if 0 <= (due - today).days <= 7 and o.get("delivery_status") != "Delivered":
            upcoming.append(o)
    pipeline = sum(o.get("total_value", 0) for o in orders if o.get("delivery_status") != "Delivered")
    stage_counts = {s: sum(1 for o in orders if o.get("current_stage") == s) for s in STAGES}
    return {
        "pipeline": pipeline,
        "outstanding": outstanding,
        "overdue": overdue,
        "upcoming": upcoming,
        "stage_counts": stage_counts,
        "delivered_this_week": sum(1 for o in orders if o.get("delivery_status") == "Delivered" and (o.get("delivery_date") or "") >= (today - timedelta(days=7)).isoformat()),
        "new_this_week": sum(1 for o in orders if (o.get("order_date") or "") >= (today - timedelta(days=7)).isoformat()),
    }

def digest_html(d):
    stages_html = "".join(f"<tr><td style='padding:6px 12px;color:#64748b;font-size:12px;'>{k}</td><td style='padding:6px 12px;color:#0f172a;font-size:13px;font-weight:600;text-align:right;'>{v}</td></tr>" for k, v in d["stage_counts"].items())
    upcoming_html = "".join(f"<tr><td style='padding:6px 12px;font-size:12px;color:#0f172a;'>{o.get('order_id')}</td><td style='padding:6px 12px;font-size:12px;color:#334155;'>{o.get('client_name')}</td><td style='padding:6px 12px;font-size:12px;color:#334155;text-align:right;'>{o.get('required_delivery_date')}</td></tr>" for o in d["upcoming"][:10])
    if not upcoming_html:
        upcoming_html = "<tr><td colspan='3' style='padding:12px;color:#94a3b8;font-size:12px;text-align:center;'>Nothing due in the next 7 days.</td></tr>"
    return f"""
    <table width='100%' cellpadding='0' cellspacing='0' style='background:#f8fafc;padding:32px 0;font-family:Arial,sans-serif;'>
      <tr><td align='center'>
        <table width='620' cellpadding='0' cellspacing='0' style='background:#fff;border-radius:10px;overflow:hidden;box-shadow:0 4px 12px rgba(15,23,42,.08);'>
          <tr><td style='background:#0f172a;padding:24px 28px;'>
            <div style='color:#93c5fd;font-size:11px;letter-spacing:2px;font-weight:700;'>PINAKI SOLUTIONS · WEEKLY DIGEST</div>
            <div style='color:#fff;font-size:22px;font-weight:700;margin-top:4px;'>Monday morning briefing</div>
            <div style='color:#cbd5e1;font-size:12px;margin-top:4px;'>{ist_now().strftime('%A, %d %b %Y')}</div>
          </td></tr>
          <tr><td style='padding:28px 28px 8px;'>
            <table cellpadding='0' cellspacing='0' style='width:100%;'>
              <tr>
                <td style='padding:14px;background:#eff6ff;border-radius:8px;width:33%;'><div style='color:#2563eb;font-size:11px;font-weight:700;letter-spacing:1px;'>OPEN PIPELINE</div><div style='color:#0f172a;font-size:22px;font-weight:800;margin-top:6px;'>₹{d["pipeline"]:,.0f}</div></td>
                <td style='width:12px;'></td>
                <td style='padding:14px;background:#fffbeb;border-radius:8px;width:33%;'><div style='color:#a16207;font-size:11px;font-weight:700;letter-spacing:1px;'>OUTSTANDING</div><div style='color:#0f172a;font-size:22px;font-weight:800;margin-top:6px;'>₹{d["outstanding"]:,.0f}</div></td>
                <td style='width:12px;'></td>
                <td style='padding:14px;background:#fef2f2;border-radius:8px;width:33%;'><div style='color:#b91c1c;font-size:11px;font-weight:700;letter-spacing:1px;'>OVERDUE</div><div style='color:#0f172a;font-size:22px;font-weight:800;margin-top:6px;'>₹{d["overdue"]:,.0f}</div></td>
              </tr>
            </table>
          </td></tr>
          <tr><td style='padding:20px 28px 8px;'>
            <div style='color:#0f172a;font-size:15px;font-weight:700;'>Due in the next 7 days</div>
            <table cellpadding='0' cellspacing='0' style='width:100%;margin-top:12px;background:#f8fafc;border-radius:8px;overflow:hidden;'>
              <tr><td style='padding:8px 12px;font-size:10px;color:#94a3b8;letter-spacing:1px;font-weight:700;'>ORDER</td><td style='padding:8px 12px;font-size:10px;color:#94a3b8;letter-spacing:1px;font-weight:700;'>CLIENT</td><td style='padding:8px 12px;font-size:10px;color:#94a3b8;letter-spacing:1px;text-align:right;font-weight:700;'>DELIVERY</td></tr>
              {upcoming_html}
            </table>
          </td></tr>
          <tr><td style='padding:20px 28px 8px;'>
            <div style='color:#0f172a;font-size:15px;font-weight:700;'>Orders by stage</div>
            <table cellpadding='0' cellspacing='0' style='width:100%;margin-top:12px;background:#f8fafc;border-radius:8px;overflow:hidden;'>
              {stages_html}
            </table>
          </td></tr>
          <tr><td style='padding:20px 28px 8px;color:#475569;font-size:13px;line-height:1.7;'>
            <b>{d["new_this_week"]}</b> new orders and <b>{d["delivered_this_week"]}</b> deliveries closed in the last 7 days.
          </td></tr>
          <tr><td style='padding:14px 28px 26px;color:#94a3b8;font-size:11px;line-height:1.6;border-top:1px solid #e2e8f0;'>
            Automated digest · Pinaki Solutions CRM
          </td></tr>
        </table>
      </td></tr>
    </table>
    """

async def send_weekly_digest(force=False):
    if not os.environ.get("RESEND_API_KEY"):
        return False
    now_ist = ist_now()
    key = now_ist.strftime("%G-W%V")  # ISO week
    state = await db.digest_state.find_one({"_id": "weekly"}) or {}
    if not force and state.get("last_week") == key:
        return False
    users = await db.users.find({"role": {"$in": ["admin", "management"]}}, {"_id": 0, "email": 1, "name": 1}).to_list(100)
    recipients = [u["email"] for u in users if u.get("email")]
    if not recipients:
        return False
    data = await build_digest()
    html = digest_html(data)
    subject = f"Pinaki weekly briefing — pipeline ₹{data['pipeline']:,.0f}, overdue ₹{data['overdue']:,.0f}"
    params = {"from": f"Pinaki Solutions <{SENDER_EMAIL}>", "to": recipients, "subject": subject, "html": html}
    try:
        result = await asyncio.to_thread(resend.Emails.send, params)
    except Exception:
        logging.exception("Weekly digest send failed")
        return False
    provider_id = result.get("id") if isinstance(result, dict) else getattr(result, "id", None)
    await db.digest_state.update_one({"_id": "weekly"}, {"$set": {"last_week": key, "last_sent_at": now(), "recipients": recipients, "provider_id": provider_id}}, upsert=True)
    logging.info(f"Weekly digest sent to {recipients} ({provider_id})")
    return True

async def digest_loop():
    await asyncio.sleep(45)
    while True:
        try:
            now_ist = ist_now()
            # Monday between 09:00 and 09:59 IST
            if now_ist.weekday() == 0 and 9 <= now_ist.hour < 10:
                await send_weekly_digest()
        except Exception:
            logging.exception("digest loop error")
        await asyncio.sleep(1800)  # every 30 minutes

@api.post("/emails/run-digest")
async def run_digest_now(user=Depends(require("admin", "management"))):
    ok = await send_weekly_digest(force=True)
    return {"ok": ok, "reason": None if ok else "RESEND_API_KEY missing or no recipients"}

@api.get("/emails/digest-preview")
async def digest_preview(user=Depends(require("admin", "management"))):
    return await build_digest()

# ---------- public approval link ----------
class PublicApprovalDecision(BaseModel):
    status: str  # Approved / Rejected / Revision Required
    approver: str
    comments: str = ""

@api.get("/public/approval/{token}")
async def public_approval_get(token: str):
    try:
        payload = verify_approval_token(token)
    except jwt.PyJWTError:
        raise HTTPException(400, "This approval link is invalid or has expired")
    order = await db.orders.find_one({"order_id": payload["order_id"]}, {"_id": 0})
    design = await db.designs.find_one({"design_id": payload["sub"]}, {"_id": 0})
    if not order or not design:
        raise HTTPException(404, "Order or design not found")
    approvals = await db.approvals.find({"order_id": order["order_id"], "design_id": design["design_id"]}, {"_id": 0}).sort("created_at", -1).to_list(10)
    return {
        "order": {k: order.get(k) for k in ("order_id", "client_name", "product", "product_code", "quantity", "unit_price", "total_value", "required_delivery_date", "current_stage", "approval_status")},
        "design": {k: design.get(k) for k in ("design_id", "version", "designer", "comments", "created_at", "file_name")},
        "documents": await db.documents.find({"order_id": order["order_id"], "category": "Design File"}, {"_id": 0}).to_list(20),
        "decided": len(approvals) > 0,
        "history": approvals,
    }

@api.post("/public/approval/{token}")
async def public_approval_post(token: str, data: PublicApprovalDecision):
    try:
        payload = verify_approval_token(token)
    except jwt.PyJWTError:
        raise HTTPException(400, "This approval link is invalid or has expired")
    if data.status not in ("Approved", "Rejected", "Revision Required"):
        raise HTTPException(400, "Invalid status")
    order = await db.orders.find_one({"order_id": payload["order_id"]}, {"_id": 0})
    if not order:
        raise HTTPException(404, "Order not found")
    approval = {
        "approval_id": new_id("AP"),
        "order_id": payload["order_id"],
        "design_id": payload["sub"],
        "status": data.status,
        "approver": data.approver,
        "comments": data.comments,
        "source": "client_link",
        "created_at": now(),
    }
    await db.approvals.insert_one(approval)
    updates = {"approval_status": data.status, "updated_at": now()}
    if data.status == "Approved":
        updates["current_stage"] = "Production"
        await notify(["production", "sales"], payload["order_id"], f"{payload['order_id']} approved by client via link")
    else:
        updates["current_stage"] = "Design"
        await notify(["design", "sales"], payload["order_id"], f"{payload['order_id']} needs {data.status.lower()} (client feedback)")
    await db.orders.update_one({"order_id": payload["order_id"]}, {"$set": updates})
    await add_history(payload["order_id"], f"Client via link: {data.status}", data.approver, data.comments)
    return {"ok": True, "status": data.status}

class InvoiceSnooze(BaseModel):
    days: int = 7
    promised_date: str = ""
    note: str = ""

@api.post("/invoices/{invoice_id}/snooze")
async def snooze_invoice(invoice_id: str, data: InvoiceSnooze, user=Depends(require("admin", "accounts"))):
    inv = await db.invoices.find_one({"invoice_id": invoice_id}, {"_id": 0})
    if not inv: raise HTTPException(404, "Invoice not found")
    if inv.get("status") == "Paid": raise HTTPException(409, "Invoice is already paid — nothing to snooze")
    days = max(1, min(int(data.days or 7), 60))
    until = (datetime.now(timezone.utc) + timedelta(days=days)).date().isoformat()
    await db.invoices.update_one({"invoice_id": invoice_id}, {"$set": {
        "snoozed_until": until,
        "snooze_promised_date": data.promised_date or None,
        "snooze_note": data.note,
        "snoozed_by": user["name"],
        "snoozed_at": now(),
    }})
    await add_history(inv["order_id"], f"Reminders snoozed until {until}" + (f" (client promised {data.promised_date})" if data.promised_date else ""), user["name"], data.note)
    return {"ok": True, "snoozed_until": until}

@api.post("/invoices/{invoice_id}/unsnooze")
async def unsnooze_invoice(invoice_id: str, user=Depends(require("admin", "accounts"))):
    inv = await db.invoices.find_one({"invoice_id": invoice_id}, {"_id": 0})
    if not inv: raise HTTPException(404, "Invoice not found")
    await db.invoices.update_one({"invoice_id": invoice_id}, {"$unset": {"snoozed_until": "", "snooze_promised_date": "", "snooze_note": "", "snoozed_by": "", "snoozed_at": ""}})
    await add_history(inv["order_id"], "Reminder snooze cleared", user["name"])
    return {"ok": True}

# ---------- daily lead follow-up reminders ----------
def followup_html(owner, leads, today):
    rows = ""
    for l in leads:
        overdue_txt = "Today" if l.get("next_follow_up") == today else f"Overdue · {l.get('next_follow_up')}"
        tone = "#dc2626" if l.get("next_follow_up") != today else "#a16207"
        rows += (
            f"<tr>"
            f"<td style='padding:10px 14px;font-size:12.5px;color:#0f172a;'><b>{l.get('company_name')}</b><br>"
            f"<span style='color:#94a3b8;font-size:11px;'>{l.get('contact_person','')} · {l.get('email','')}</span></td>"
            f"<td style='padding:10px 14px;font-size:11px;color:#334155;'>{l.get('status')}</td>"
            f"<td style='padding:10px 14px;font-size:11px;color:{tone};text-align:right;font-weight:700;'>{overdue_txt}</td>"
            f"</tr>"
        )
    return f"""
    <table width='100%' cellpadding='0' cellspacing='0' style='background:#f8fafc;padding:32px 0;font-family:Arial,sans-serif;'>
      <tr><td align='center'>
        <table width='600' cellpadding='0' cellspacing='0' style='background:#fff;border-radius:10px;overflow:hidden;box-shadow:0 4px 12px rgba(15,23,42,.08);'>
          <tr><td style='background:#0f172a;padding:22px 28px;'>
            <div style='color:#93c5fd;font-size:11px;letter-spacing:2px;font-weight:700;'>PINAKI SOLUTIONS · SALES</div>
            <div style='color:#fff;font-size:22px;font-weight:700;margin-top:4px;'>Today's follow-ups</div>
            <div style='color:#cbd5e1;font-size:12px;margin-top:2px;'>Hi {owner}, {len(leads)} lead(s) need a touch today.</div>
          </td></tr>
          <tr><td style='padding:14px 28px 22px;'>
            <table cellpadding='0' cellspacing='0' style='width:100%;background:#f8fafc;border-radius:8px;overflow:hidden;'>
              <tr><td style='padding:8px 14px;font-size:10px;color:#94a3b8;letter-spacing:1px;font-weight:700;'>LEAD</td><td style='padding:8px 14px;font-size:10px;color:#94a3b8;letter-spacing:1px;font-weight:700;'>STATUS</td><td style='padding:8px 14px;font-size:10px;color:#94a3b8;letter-spacing:1px;text-align:right;font-weight:700;'>DUE</td></tr>
              {rows}
            </table>
          </td></tr>
          <tr><td style='padding:12px 28px 22px;color:#94a3b8;font-size:11px;border-top:1px solid #e2e8f0;'>Automated daily digest · Pinaki Solutions CRM</td></tr>
        </table>
      </td></tr>
    </table>
    """

async def send_followup_reminders(force=False):
    now_ist_dt = ist_now()
    key = now_ist_dt.strftime("%Y-%m-%d")
    state = await db.followup_state.find_one({"_id": "daily"}) or {}
    if not force and state.get("last_date") == key:
        return {"sent": 0, "reason": "already sent today", "leads": 0}
    today = now_ist_dt.date().isoformat()
    leads = await db.leads.find({
        "status": {"$nin": ["Won", "Lost"]},
        "next_follow_up": {"$lte": today, "$ne": ""},
    }, {"_id": 0}).to_list(2000)
    by_owner = {}
    for l in leads:
        by_owner.setdefault(l.get("assigned_to") or "Sales", []).append(l)
    email_sent = 0
    for owner, lst in by_owner.items():
        # in-app notifications for every lead
        for lead in lst:
            await db.notifications.insert_one({
                "notif_id": new_id("NT"), "role": "sales", "order_id": lead.get("lead_id"),
                "message": f"Follow up with {lead['company_name']} — {'today' if lead['next_follow_up']==today else 'overdue since '+lead['next_follow_up']}",
                "read": False, "created_at": now(),
            })
        # one consolidated email to the owner
        if os.environ.get("RESEND_API_KEY"):
            user = await db.users.find_one({"name": owner}, {"_id": 0, "email": 1, "name": 1})
            if user and user.get("email"):
                subject = f"Today's follow-ups — {len(lst)} lead(s) need attention"
                params = {"from": f"Pinaki Solutions <{SENDER_EMAIL}>", "to": [user["email"]], "subject": subject, "html": followup_html(owner, lst, today)}
                try:
                    result = await asyncio.to_thread(resend.Emails.send, params)
                    provider_id = result.get("id") if isinstance(result, dict) else getattr(result, "id", None)
                    await db.emails.insert_one({"email_id": new_id("EM"), "order_id": None, "client_name": None, "recipient": user["email"], "cc": [], "subject": subject, "template": "followup_daily", "provider_id": provider_id, "sent_by": "Automation", "created_at": now(), "auto": True})
                    email_sent += 1
                except Exception:
                    logging.exception("followup email send failed")
    await db.followup_state.update_one({"_id": "daily"}, {"$set": {"last_date": key, "last_run_at": now(), "leads": len(leads), "emails_sent": email_sent}}, upsert=True)
    return {"sent": email_sent, "leads": len(leads), "owners": len(by_owner)}

async def followup_loop():
    await asyncio.sleep(60)
    while True:
        try:
            now_ist_dt = ist_now()
            if 9 <= now_ist_dt.hour < 10:
                await send_followup_reminders()
        except Exception:
            logging.exception("followup loop error")
        await asyncio.sleep(1800)

@api.post("/emails/run-followups")
async def run_followups_now(user=Depends(require("admin", "sales", "management"))):
    return await send_followup_reminders(force=True)

# ---------- background: overdue reminders ----------
async def send_overdue_reminder(invoice, order, days_late):
    if not os.environ.get("RESEND_API_KEY"):
        return False
    client = await db.clients.find_one({"client_id": invoice.get("client_id")}, {"_id": 0})
    recipient = (client or {}).get("email")
    if not recipient:
        return False
    tpl = EMAIL_TEMPLATES["payment_reminder"]
    subject = tpl["subject"].format(order_id=order.get("order_id"), invoice_number=invoice["invoice_number"]) + f" ({days_late} day{'s' if days_late != 1 else ''} overdue)"
    body = (
        f"This is an automated reminder — invoice {invoice['invoice_number']} is now "
        f"{days_late} days past its due date of {invoice['due_date']}. "
        f"Please arrange payment at your earliest convenience or reply to this email if there is a query."
    )
    attachments = []
    if invoice.get("pdf_document_id"):
        d = await db.documents.find_one({"document_id": invoice["pdf_document_id"]}, {"_id": 0})
        if d:
            try:
                stream = await gridfs.open_download_stream(ObjectId(d["storage_id"]))
                content = await stream.read()
                attachments.append({"filename": d["name"], "content": base64.b64encode(content).decode()})
            except Exception:
                pass
    html = build_email_html("payment_reminder", order, body, "Pinaki Accounts", invoice, pay_url=f"{os.environ.get('FRONTEND_URL', '')}/pay/{sign_pay_token(invoice['invoice_id'])}" if os.environ.get("FRONTEND_URL") else None)
    params = {"from": f"Pinaki Solutions <{SENDER_EMAIL}>", "to": [recipient], "subject": subject, "html": html}
    if attachments:
        params["attachments"] = attachments
    try:
        result = await asyncio.to_thread(resend.Emails.send, params)
    except Exception:
        logging.exception("Auto reminder send failed")
        return False
    provider_id = result.get("id") if isinstance(result, dict) else getattr(result, "id", None)
    await db.emails.insert_one({
        "email_id": new_id("EM"), "order_id": order["order_id"], "client_name": order.get("client_name"),
        "recipient": recipient, "cc": [], "subject": subject, "template": "payment_reminder",
        "provider_id": provider_id, "document_ids": [invoice.get("pdf_document_id")] if invoice.get("pdf_document_id") else [],
        "sent_by": "Automation", "created_at": now(), "auto": True, "days_late": days_late,
    })
    await db.invoices.update_one({"invoice_id": invoice["invoice_id"]}, {"$addToSet": {"reminders_sent": days_late}})
    await add_history(order["order_id"], f"Auto reminder sent — {days_late} days overdue", "Automation")
    return True

async def send_overdue_escalation(invoice, order, days_late):
    if not os.environ.get("RESEND_API_KEY"):
        return False
    # gather internal recipients: sales owner + management + admins
    recipients = set()
    mgmt = await db.users.find({"role": {"$in": ["management", "admin"]}}, {"_id": 0, "email": 1}).to_list(50)
    for u in mgmt:
        if u.get("email"):
            recipients.add(u["email"])
    sales_owner_name = order.get("sales_owner")
    if sales_owner_name:
        sales_user = await db.users.find_one({"name": sales_owner_name}, {"_id": 0, "email": 1})
        if sales_user and sales_user.get("email"):
            recipients.add(sales_user["email"])
    if not recipients:
        return False
    client = await db.clients.find_one({"client_id": invoice.get("client_id")}, {"_id": 0})
    prior_reminders = [r for r in invoice.get("reminders_sent", []) if isinstance(r, int)]
    reminder_line = f"Auto reminders already sent at {sorted(prior_reminders)} days late." if prior_reminders else "No auto reminders sent yet."
    subject = f"ESCALATION · {invoice['invoice_number']} · {days_late} days overdue · {order.get('client_name')}"
    body = (
        f"Invoice {invoice['invoice_number']} for {order.get('client_name')} "
        f"is now {days_late} days past its due date of {invoice['due_date']} and remains unpaid. "
        f"{reminder_line} Please intervene directly.\n\n"
        f"Client contact: {(client or {}).get('contact_person') or '—'} · {(client or {}).get('email') or '—'} · {(client or {}).get('phone') or '—'}\n"
        f"Payment terms: {invoice.get('payment_terms', '—')}"
    )
    attachments = []
    if invoice.get("pdf_document_id"):
        d = await db.documents.find_one({"document_id": invoice["pdf_document_id"]}, {"_id": 0})
        if d:
            try:
                stream = await gridfs.open_download_stream(ObjectId(d["storage_id"]))
                content = await stream.read()
                attachments.append({"filename": d["name"], "content": base64.b64encode(content).decode()})
            except Exception:
                pass
    html = build_email_html("payment_reminder", order, body, "Pinaki Automation", invoice)
    # inject a red escalation banner at the top of the HTML
    html = html.replace(
        "<div style='color:#fff;font-size:22px;font-weight:700;margin-top:4px;'>Order lifecycle update</div>",
        f"<div style='background:#dc2626;color:#fff;display:inline-block;padding:4px 10px;border-radius:4px;font-size:10px;letter-spacing:1.4px;font-weight:800;margin-bottom:8px;'>INTERNAL ESCALATION · {days_late}D OVERDUE</div><div style='color:#fff;font-size:22px;font-weight:700;'>Payment escalation</div>",
        1,
    )
    params = {"from": f"Pinaki Solutions <{SENDER_EMAIL}>", "to": sorted(recipients), "subject": subject, "html": html}
    if attachments:
        params["attachments"] = attachments
    try:
        result = await asyncio.to_thread(resend.Emails.send, params)
    except Exception:
        logging.exception("Escalation send failed")
        return False
    provider_id = result.get("id") if isinstance(result, dict) else getattr(result, "id", None)
    await db.emails.insert_one({
        "email_id": new_id("EM"), "order_id": order["order_id"], "client_name": order.get("client_name"),
        "recipient": ", ".join(sorted(recipients)), "cc": [], "subject": subject, "template": "escalation",
        "provider_id": provider_id, "document_ids": [invoice.get("pdf_document_id")] if invoice.get("pdf_document_id") else [],
        "sent_by": "Automation", "created_at": now(), "auto": True, "days_late": days_late, "escalation": True,
    })
    await db.invoices.update_one({"invoice_id": invoice["invoice_id"]}, {"$addToSet": {"reminders_sent": "escalation_14"}})
    await notify(["management", "sales", "accounts"], order["order_id"], f"Escalation sent — {invoice['invoice_number']} {days_late}d overdue")
    await add_history(order["order_id"], f"Overdue escalation sent to sales + management ({days_late}d overdue)", "Automation")
    return True

async def check_overdue_reminders():
    invoices = await db.invoices.find({"status": {"$nin": ["Paid"]}}, {"_id": 0}).to_list(1000)
    payments = await db.payments.find({}, {"_id": 0}).to_list(2000)
    paid_map = {}
    for p in payments:
        paid_map[p["invoice_id"]] = paid_map.get(p["invoice_id"], 0) + p.get("amount", 0)
    today_dt = datetime.now(timezone.utc).date()
    for inv in invoices:
        if paid_map.get(inv["invoice_id"], 0) >= inv.get("total", 0):
            continue
        # respect snooze
        snoozed_until = inv.get("snoozed_until")
        if snoozed_until:
            try:
                if datetime.fromisoformat(snoozed_until).date() >= today_dt:
                    continue
            except Exception:
                pass
        try:
            due = datetime.fromisoformat(inv["due_date"]).date()
        except Exception:
            continue
        days_late = (today_dt - due).days
        already = inv.get("reminders_sent", [])
        if days_late <= 0:
            continue
        order = await db.orders.find_one({"order_id": inv["order_id"]}, {"_id": 0})
        if not order:
            continue
        # 14+ day escalation to internal team
        if days_late >= 14 and "escalation_14" not in already:
            await send_overdue_escalation(inv, order, days_late)
            continue
        # 3 / 7 day client reminders
        for threshold in (3, 7):
            if days_late >= threshold and threshold not in already:
                await send_overdue_reminder(inv, order, days_late)
                break

async def reminder_loop():
    # small delay so it doesn't hammer during startup
    await asyncio.sleep(30)
    while True:
        try:
            await check_overdue_reminders()
        except Exception:
            logging.exception("reminder loop iteration failed")
        await asyncio.sleep(3600)  # once per hour

@api.post("/emails/run-reminders")
async def run_reminders_now(user=Depends(require("admin", "accounts"))):
    await check_overdue_reminders()
    return {"ok": True}

# ---------- users ----------
@api.get("/users")
async def list_users(user=Depends(require("admin", "management"))):
    return await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(100)

# ---------- seed ----------
SEED_ORDERS = [
    {"client_id": "CL-1001", "product": "Retail display kits", "product_code": "RDK-A", "quantity": 1200, "unit_price": 285, "required_delivery_date": "2026-04-08", "priority": "High", "current_stage": "Production", "approval_status": "Approved", "production_status": "In Progress", "po_number": "PO/NS/2026/0311"},
    {"client_id": "CL-1002", "product": "Branded menu folders", "product_code": "MEN-STD", "quantity": 500, "unit_price": 180, "required_delivery_date": "2026-04-02", "priority": "Normal", "current_stage": "Client Approval", "approval_status": "Pending Approval", "production_status": "Not Started", "po_number": "AST/26/007"},
    {"client_id": "CL-1003", "product": "Trade show backdrop", "product_code": "TSB-4X8", "quantity": 4, "unit_price": 14500, "required_delivery_date": "2026-03-30", "priority": "Urgent", "current_stage": "Delivery", "approval_status": "Approved", "production_status": "Completed", "dispatch_status": "Dispatched", "delivery_status": "In Transit", "po_number": "SVI-PO-1188"},
    {"client_id": "CL-1001", "product": "Corrugated shipper boxes", "product_code": "CSB-M", "quantity": 5000, "unit_price": 42, "required_delivery_date": "2026-03-18", "priority": "Normal", "current_stage": "Payment", "approval_status": "Approved", "production_status": "Completed", "dispatch_status": "Dispatched", "delivery_status": "Delivered", "invoice_status": "Generated", "po_number": "PO/NS/2026/0288"},
]

@app.on_event("startup")
async def seed():
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("email", unique=True)
    await db.orders.create_index("order_id", unique=True)
    await db.invoices.create_index("invoice_number", unique=True)
    await db.challans.create_index("challan_number", unique=True)
    users = [
        ("admin@pinakisolutions.com", "Admin", "admin"),
        ("sales@pinakisolutions.com", "Riya Shah", "sales"),
        ("design@pinakisolutions.com", "Kabir Rao", "design"),
        ("production@pinakisolutions.com", "Neha Joshi", "production"),
        ("dispatch@pinakisolutions.com", "Rohit Iyer", "dispatch"),
        ("accounts@pinakisolutions.com", "Vikram Singh", "accounts"),
        ("management@pinakisolutions.com", "Priya Nair", "management"),
    ]
    for email, name, role in users:
        if not await db.users.find_one({"email": email}):
            await db.users.insert_one({"user_id": f"USR-{secrets.token_hex(4)}", "email": email, "name": name, "role": role, "password_hash": hash_password("Pinaki@123"), "created_at": now()})
    demo_clients = [
        {"client_id": "CL-1001", "company_name": "Northstar Retail Group", "contact_person": "Meera Kapoor", "email": "meera@northstar.example", "phone": "+91 98200 11882", "gstin": "27AAACN1234A1ZP", "pan": "AAACN1234A", "payment_terms": "Net 30", "credit_limit": 2500000, "account_manager": "Riya Shah", "status": "Active", "billing_address": "12 Cuffe Parade, Mumbai 400005", "shipping_address": "Northstar DC, Bhiwandi 421302", "notes": "Priority partner since 2021", "created_at": now()},
        {"client_id": "CL-1002", "company_name": "Aster Hospitality", "contact_person": "Dev Malhotra", "email": "dev@aster.example", "phone": "+91 98111 22004", "gstin": "27AABCA9876B1ZX", "pan": "AABCA9876B", "payment_terms": "50% Advance", "credit_limit": 1000000, "account_manager": "Riya Shah", "status": "Active", "billing_address": "8 MG Road, Bengaluru 560001", "shipping_address": "Aster Central Warehouse, Hosur", "notes": "Boutique hospitality group", "created_at": now()},
        {"client_id": "CL-1003", "company_name": "Svelte Interiors", "contact_person": "Ananya Bose", "email": "ananya@svelte.example", "phone": "+91 90080 44112", "gstin": "27AAECS4433B1ZH", "pan": "AAECS4433B", "payment_terms": "Net 45", "credit_limit": 1500000, "account_manager": "Riya Shah", "status": "Active", "billing_address": "3rd Floor, Pheonix Marketcity, Pune", "shipping_address": "Svelte Studio, Kharadi, Pune", "notes": "Project-based orders", "created_at": now()},
    ]
    for c in demo_clients:
        await db.clients.update_one({"client_id": c["client_id"]}, {"$setOnInsert": c}, upsert=True)
    if await db.leads.count_documents({}) == 0:
        demo_leads = [
            {"lead_id": "LD-1001", "company_name": "Zenith Cafe Co.", "contact_person": "Aarav Mehta", "email": "aarav@zenithcafe.example", "phone": "+91 98212 33440", "source": "Referral", "status": "Qualified", "estimated_value": 350000, "product_interest": "Custom coffee bag printing", "assigned_to": "Riya Shah", "next_follow_up": "2026-02-20", "notes": "Chain of 12 cafes, wants sustainable packaging", "activities": [{"type": "call", "summary": "Intro call — very interested", "by": "Riya Shah", "at": now()}, {"type": "status_change", "summary": "Moved from Contacted to Qualified", "by": "Riya Shah", "at": now()}], "created_by": "Riya Shah", "created_at": now()},
            {"lead_id": "LD-1002", "company_name": "Lumen Studios", "contact_person": "Sara Iyer", "email": "sara@lumen.example", "phone": "+91 90099 54123", "source": "Website", "status": "Proposal Sent", "estimated_value": 850000, "product_interest": "Retail store signage kit", "assigned_to": "Riya Shah", "next_follow_up": "2026-02-18", "notes": "Awaiting feedback on the ₹8.5L proposal", "activities": [{"type": "email", "summary": "Proposal V1 sent", "by": "Riya Shah", "at": now()}], "created_by": "Riya Shah", "created_at": now()},
            {"lead_id": "LD-1003", "company_name": "Karvan Logistics", "contact_person": "Ishaan Verma", "email": "ishaan@karvan.example", "phone": "+91 98999 22014", "source": "Cold Call", "status": "Contacted", "estimated_value": 180000, "product_interest": "Corrugated shipper boxes (bulk)", "assigned_to": "Riya Shah", "next_follow_up": "2026-02-22", "notes": "Sent samples; awaiting confirmation on volume", "activities": [{"type": "call", "summary": "Discovery call done", "by": "Riya Shah", "at": now()}], "created_by": "Riya Shah", "created_at": now()},
            {"lead_id": "LD-1004", "company_name": "Bloom & Bean", "contact_person": "Nikhil Kapoor", "email": "nikhil@bloombean.example", "phone": "+91 87800 99127", "source": "LinkedIn", "status": "New", "estimated_value": 90000, "product_interest": "Menu folders + table tents", "assigned_to": "Riya Shah", "next_follow_up": "2026-02-19", "notes": "Fresh inquiry from LinkedIn ad", "activities": [{"type": "note", "summary": "Lead created", "by": "Riya Shah", "at": now()}], "created_by": "Riya Shah", "created_at": now()},
        ]
        await db.leads.insert_many(demo_leads)
    if await db.orders.count_documents({}) == 0:
        clients_map = {c["client_id"]: c["company_name"] async for c in db.clients.find({}, {"_id": 0})}
        for i, o in enumerate(SEED_ORDERS):
            subtotal = o["quantity"] * o["unit_price"]
            gst = round(subtotal * 0.18, 2)
            base = {
                "order_id": f"PS-2603-{secrets.token_hex(2).upper()}",
                "order_date": today(),
                "client_name": clients_map.get(o["client_id"], "Client"),
                "order_value": subtotal,
                "gst": gst,
                "total_value": subtotal + gst,
                "outstanding_amount": subtotal + gst,
                "sales_owner": "Riya Shah",
                "dispatch_status": "Pending",
                "delivery_status": "Pending",
                "invoice_status": "Pending",
                "approval_status": "Pending Approval",
                "production_status": "Not Started",
                "history": [{"event": "Order created", "by": "Riya Shah", "at": now()}],
                "created_at": now(),
            }
            base.update(o)
            await db.orders.insert_one(base)
            if o.get("invoice_status") == "Generated":
                inv_no = f"INV-2603-{secrets.token_hex(2).upper()}"
                await db.invoices.insert_one({"invoice_id": new_id("IV"), "invoice_number": inv_no, "order_id": base["order_id"], "client_id": o["client_id"], "client_name": base["client_name"], "invoice_date": today(), "due_date": "2026-04-20", "payment_terms": "Net 30", "taxable": subtotal, "gst": gst, "total": subtotal + gst, "status": "Generated", "created_at": now()})
                await db.orders.update_one({"order_id": base["order_id"]}, {"$set": {"invoice_number": inv_no, "current_stage": "Payment"}})
    # start background reminder loop
    asyncio.create_task(reminder_loop())
    asyncio.create_task(digest_loop())
    asyncio.create_task(followup_loop())

app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origin_regex=r"https://.*|http://localhost(:\d+)?",
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"],
)
logging.basicConfig(level=logging.INFO)

@app.on_event("shutdown")
async def shutdown():
    client.close()
